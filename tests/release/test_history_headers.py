import pathlib,sys,tempfile,unittest
import openpyxl
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[2]/'scripts/imports'))
from history import prepare
from common import Invalid
class Headers(unittest.TestCase):
 def test_each_header_position_preserves_first_real_row_and_name(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=pathlib.Path(tmp)/'synthetic.xlsx';w=openpyxl.Workbook();w.remove(w.active)
   for index,header in enumerate([3,4],1):
    s=w.create_sheet('Exact worksheet '+str(index)+' ')
    for _ in range(header-1):s.append(['title'])
    s.append(['Date','Student ID','Student Name','Applied','Attended','Passed','School'])
    s.append([None,'S0','Example Student']);s.append([None,'S1','Synthetic candidate'])
   w.save(p);result=prepare({'morarji':p},[],b'k'*32)
   self.assertEqual(len(result['held']),2);self.assertEqual(result['accepted'],[])
   sheets=result['files'][0]['sheets'] if 'files' in result else result['source_files'][0]['sheets']
   self.assertEqual([s['header_row'] for s in sheets],[3,4]);self.assertEqual([s['rows'][0]['row_number'] for s in sheets],[5,6]);self.assertEqual(sheets[0]['original_name'],'Exact worksheet 1 ')
 def test_unknown_header_is_not_guessed(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=pathlib.Path(tmp)/'synthetic.xlsx';w=openpyxl.Workbook();w.active.append(['no approved header']);w.save(p)
   with self.assertRaises(Invalid):prepare({'morarji':p},[],b'k'*32)
if __name__=='__main__':unittest.main()
