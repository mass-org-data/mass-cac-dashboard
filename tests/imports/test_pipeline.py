import sys,os,json,copy,tempfile,unittest,subprocess
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'scripts/imports'))
from common import Invalid,digest,ident,canonical,write_private
from stage import stage
from review import scaffold,validate,preview
from render import render
from adapters import observation
SHEETS=('Mpokket','Fanuc','McAfee','Janashkati Foundation','Dasar UK')
def fixture(folder):
 path=Path(folder)/'synthetic-directory.xlsx';book=Workbook();book.remove(book.active);layout={'format':'mass-source-layout-v1','kind':'directory','mapping_version':'synthetic-v1','sheets':{}}
 fields=['donor','centre','region','teacher','officer','shift','student_name']
 for n,sheet in enumerate(SHEETS,1):
  s=book.create_sheet(sheet);s.append(fields);s.append([f'Synthetic donor {n}',f'Synthetic village {n}','Synthetic region',f'Synthetic teacher {n}',f'Synthetic officer {n}','Morning','Unmapped synthetic learner'])
  layout['sheets'][sheet]={'header_row':1,'header_version':1,'columns':{k:chr(65+i) for i,k in enumerate(fields)},'verified_headers':{k:k for k in fields}}
 book.save(path);path.chmod(0o600)
 plan=stage(path,layout,b'x'*32);mapping=scaffold(plan)
 for i,o in enumerate(plan['observations'],1):
  row=o['source_row_id'];mapping['dispositions'][row]='quarantined' if i==5 else 'accepted'
  if i==5:continue
  if o['issues']:mapping['resolutions'][row]={'resolved_issues':list(o['issues']),'decision_ref':'synthetic-reviewed'}
  donor=ident('approved-donor',i);centre=ident('approved-centre',i)
  mapping['entities']['donors'].append({'id':donor,'label':f"Synthetic donor {i} O'Neil $mass_import$",'source_label':o['raw']['donor'],'source_row_id':row,'source_slot':1})
  mapping['entities']['centres'].append({'id':centre,'label':o['raw']['centre'],'donor_id':donor,'source_row_id':row,'source_slot':2})
  mapping['row_identities'][row]={'centre_id':centre,'donor_id':donor,'teacher_id':ident('approved-staff',i,'Teacher'),'officer_id':ident('approved-staff',i,'Cluster Officer')}
  for role,col in [('Teacher',4),('Cluster Officer',5)]:
   staff=ident('approved-staff',i,role);raw=o['raw']['teacher' if role=='Teacher' else 'officer']
   mapping['entities']['staff_members'].append({'id':staff,'display_name':raw,'source_name':raw,'source_row_id':row,'source_slot':col})
   mapping['entities']['centre_staff_assignments'].append({'id':ident('approved-assignment',i,role),'centre_id':centre,'staff_member_id':staff,'role_code':role,'shift_code':'Morning','source_row_id':row,'source_slot':col})
 approval={'approved':True,'scope':'directory-only','reviewer_ref':'synthetic-reviewer','decision_ref':'synthetic-decision',**{k:True for k in ('identity_reviewed','source_provenance_reviewed','single_donor_reviewed','regions_reviewed','shifts_reviewed','directory_privacy_approved')}}
 def approved():return {**approval,'plan_digest':digest(plan),'mapping_digest':digest(mapping)}
 return path,layout,plan,mapping,approved
class PipelineTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
  self.path,self.layout,self.plan,self.mapping,self.approved=fixture(self.tmp.name)
 def test_source_ids_and_keyed_rows_are_repeatable(self):
  self.assertEqual(stage(self.path,self.layout,b'x'*32),self.plan)
  self.assertEqual([s['original_name'] for s in self.plan['sheets']],list(SHEETS))
  self.assertNotEqual(stage(self.path,self.layout,b'y'*32)['sheets'][0]['rows'][1]['content_hash'],self.plan['sheets'][0]['rows'][1]['content_hash'])
 def test_approved_preview_has_explicit_directory_not_student_counts(self):
  self.assertEqual(validate(self.plan,self.mapping,self.approved())['proposed_entities'],dict(donors=4,centres=4,staff_members=8,centre_staff_assignments=8))
  self.assertFalse(preview(self.plan,self.mapping)['approved_for_execution'])
 def test_stale_review_digest_rejected(self):
  a=self.approved();self.mapping['entities']['donors'][0]['label']='Changed'
  with self.assertRaises(Invalid):validate(self.plan,self.mapping,a)
 def test_no_automatic_identity_or_privacy_approval(self):
  a=self.approved();a['directory_privacy_approved']=False
  with self.assertRaises(Invalid):render(self.plan,self.mapping,a)
  with self.assertRaises(Invalid):validate(self.plan,scaffold(self.plan),self.approved())
 def test_canonical_duplicate_or_source_collision_rejected(self):
  self.mapping['entities']['staff_members'].append(copy.deepcopy(self.mapping['entities']['staff_members'][0]))
  with self.assertRaises(Invalid):validate(self.plan,self.mapping,self.approved())
 def test_unresolved_region_identity_blocks_accepted_row(self):
  self.plan['observations'][0]['issues'].append('masaraguppi_review');self.mapping['plan_digest']=digest(self.plan)
  with self.assertRaises(Invalid):validate(self.plan,self.mapping,self.approved())
 def test_inferred_shift_and_missing_parent_rejected(self):
  a=self.mapping['entities']['centre_staff_assignments'][0];a['shift_code']='Evening'
  with self.assertRaises(Invalid):validate(self.plan,self.mapping,self.approved())
  a['shift_code']='Morning';a['centre_id']=ident('unknown')
  with self.assertRaises(Invalid):validate(self.plan,self.mapping,self.approved())
 def test_sql_escaping_and_minimized_payload(self):
  sql=render(self.plan,self.mapping,self.approved())
  self.assertIn("O''Neil",sql);self.assertNotIn('Unmapped synthetic learner',sql)
  self.assertNotIn('original_name',sql);self.assertNotIn('fingerprint_key_id',sql)
  self.assertNotIn('CREATE TABLE',sql);self.assertNotIn('GRANT ',sql)
  self.assertRegex(sql,r'DO \$mass_import_[0-9a-f]+\$')
  import re
  delimiter=re.search(r'DO (\$mass_import_[0-9a-f]+\$)',sql).group(1);self.assertEqual(sql.count(delimiter),2)
  self.assertIn("O''Neil $mass_import$",sql)
 def test_sensitive_or_drifting_layout_rejected(self):
  self.layout['sheets']['Mpokket']['columns']['phone']='H'
  with self.assertRaises(Invalid):stage(self.path,self.layout,b'x'*32)
 def test_private_artifact_permissions_and_no_overwrite(self):
  out=Path(self.tmp.name)/'private.json';write_private(out,'{}');self.assertEqual(out.stat().st_mode&0o777,0o600)
  with self.assertRaises(FileExistsError):write_private(out,'changed')
  with self.assertRaises(Invalid):write_private(ROOT/'unsafe-private.json','{}')
 def test_historical_adapters_do_not_infer_missing_business_facts(self):
  for kind in ('child','attendance','assessment','activity','morarji','visit'):
   o=observation(kind,{})
   self.assertFalse(o['operational_import_ready']);self.assertIn('unverified_identity',o['issues'])
  self.assertIn('unrecorded_unknown',observation('attendance',{},dict(date=datetime(2026,4,1),code=''))['issues'])
  self.assertEqual(observation('attendance',{},dict(date=datetime(2026,4,1),code='P'))['values']['attendance_date'],'2026-04-01')
  self.assertIn('historical_scale_conflict_no_rescale',observation('assessment',dict(score=19,max_score=100))['issues'])
 def test_historical_sql_import_is_not_silently_enabled(self):
  self.plan['kind']='assessment';self.mapping['plan_digest']=digest(self.plan)
  with self.assertRaises(Invalid):render(self.plan,self.mapping,self.approved())
 def test_cli_private_stage_preview_and_unapproved_template_fail_closed(self):
  directory=Path(self.tmp.name)
  layout=directory/'layout.json';write_private(layout,canonical(self.layout))
  cli=[sys.executable,str(ROOT/'scripts/imports/pipeline.py')]
  def run(*args):return subprocess.run(cli+list(args),capture_output=True,text=True)
  key=directory/'key';self.assertEqual(run('init-key','--out',str(key)).returncode,0);self.assertEqual(key.stat().st_mode&0o777,0o600)
  plan=directory/'plan.json';r=run('stage','--workbook',str(self.path),'--layout',str(layout),'--key',str(key),'--out',str(plan))
  self.assertEqual(r.returncode,0);self.assertNotIn('Synthetic',r.stdout)
  mapping=directory/'mapping.json';self.assertEqual(run('review-template','--plan',str(plan),'--out',str(mapping)).returncode,0)
  report=directory/'preview.json';self.assertEqual(run('preview','--plan',str(plan),'--mapping',str(mapping),'--out',str(report)).returncode,0)
  self.assertEqual(json.loads(report.read_text())['populated_student_rows'],5)
  approval=directory/'approval.json';self.assertEqual(run('approval-template','--plan',str(plan),'--mapping',str(mapping),'--out',str(approval)).returncode,0)
  self.assertFalse(json.loads(approval.read_text())['approved'])
  sql=directory/'denied.sql';r=run('render','--plan',str(plan),'--mapping',str(mapping),'--approval',str(approval),'--workbook',str(self.path),'--out',str(sql))
  self.assertEqual(r.returncode,2);self.assertFalse(sql.exists());self.assertNotIn('Synthetic',r.stderr)
 def test_all_six_historical_workbook_families_stage_with_explicit_layouts(self):
  for kind,raw in [('child',dict(student_name='Synthetic',enrolled=None)),('attendance',dict(student_name='Synthetic')),('assessment',dict(student_name='Synthetic',score=0,max_score=25,assessed_on=datetime(2026,4,1),academic_year=None)),('activity',dict(student_name='Synthetic',month='April',slot_1='Yes')),('morarji',dict(student_name='Synthetic',passed='Yes',attended='No')),('visit',dict(officer='Synthetic',visited_on=None))]:
   book=Workbook();s=book.active;s.title='Exact tab ';fields=list(raw);s.append(fields);s.append(list(raw.values()))
   cfg={'header_row':1,'columns':{k:chr(65+i) for i,k in enumerate(fields)},'verified_headers':{k:k for k in fields}}
   if kind=='attendance':s.cell(1,3,datetime(2026,4,1));s.cell(2,3,'P');cfg['attendance_columns']={'C':'2026-04-01T00:00:00'}
   p=Path(self.tmp.name)/(kind+'.xlsx');book.save(p);p.chmod(0o600)
   result=stage(p,{'format':'mass-source-layout-v1','kind':kind,'mapping_version':'synthetic-v1','sheets':{'Exact tab ':cfg}},b'x'*32)
   self.assertEqual(result['sheets'][0]['original_name'],'Exact tab ');self.assertEqual(len(result['observations']),1)
   if kind=='assessment':self.assertEqual(result['observations'][0]['values']['verified_date'],'2026-04-01')
if __name__=='__main__':unittest.main()
