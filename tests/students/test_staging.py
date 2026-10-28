"""Synthetic source/staging invariants; no actual MASS values."""
import copy,hashlib,json,pathlib,sys,tempfile,unittest,uuid
from datetime import datetime
import openpyxl
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[2]/'scripts/imports'))
from students import prepare,render,HEADERS
from history import prepare as historical
from common import Invalid

class Staging(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.p=pathlib.Path(self.tmp.name)
        self.source=self.p/'source.xlsx';self.staging=self.p/'staging.xlsx'
        w=openpyxl.Workbook();s=w.active;s.title='Exact Sheet '
        for i in range(3):s.append(['heading'])
        s.append(['Synthetic Donor','Synthetic learner','B Synthetic','Female','Grade 1',None,'Morning']);w.save(self.source)
        self.id=str(uuid.uuid4());self.source_id=str(uuid.uuid4());self.centre_id=str(uuid.uuid4())
        self.row=['ST_BS_1_01',self.id,'Synthetic Donor','B Synthetic','Synthetic learner','Grade 1','Morning','READY_FOR_DATABASE_MATCH','Exact Sheet ',4,'2026-04-01','2027-04-01']
        self.plan={'source':{'sha256':hashlib.sha256(self.source.read_bytes()).hexdigest()},'sheets':[{'original_name':'Exact Sheet ','ordinal':1,'rows':[{'row_number':4,'content_hash':'a'*64,'id':self.source_id}]}]}
        self.live=[{'sheet_ordinal':1,'row_number':4,'content_hash':'a'*64,'source_row_id':self.source_id,'sha256':self.plan['source']['sha256']}]
        self.centres=[{'donor_label':'Synthetic Donor','centre_label':'B Synthetic','centre_id':self.centre_id}]
    def stage(self,rows=None):
        w=openpyxl.Workbook();s=w.active;s.title='Student import staging';s.append(HEADERS)
        for r in rows or [self.row]:s.append(r)
        w.save(self.staging)
        return prepare(self.staging,self.source,self.plan,self.centres,self.live,b'k'*32)
    def test_ready_checks_source_and_keeps_ids(self):
        p=self.stage();self.assertEqual(p['counts'],{'source':1,'accepted':1,'held':0,'held_reasons':{}})
        r=p['rows'][0];self.assertEqual((r['child_id'],r['student_id'],r['gender']),(self.id,'ST_BS_1_01','Girl'))
    def test_missing_shift_is_held_not_fabricated(self):
        self.row[6]=None;self.row[7]='MISSING_SHIFT';p=self.stage();self.assertEqual(p['rows'],[]);self.assertEqual(p['counts']['held'],1)
    def test_duplicate_physical_source_rejected(self):
        with self.assertRaises(Invalid):self.stage([self.row,self.row])
    def test_modified_source_name_rejected(self):
        self.row[4]='Synthetic other learner'
        with self.assertRaises(Invalid):self.stage()
    def test_missing_centre_held(self):
        self.centres=[];self.assertEqual(self.stage()['held'][0]['reason'],'ambiguous_or_missing_centre')
    def test_source_fingerprint_drift_rejected(self):
        self.live[0]['content_hash']='b'*64
        with self.assertRaises(Invalid):self.stage()
    def test_unknown_shift_does_not_gain_ready_status(self):
        self.row[6]='Both';self.assertIn('unverified_shift',self.stage()['held'][0]['reason'])
    def test_assumed_dates_are_explicit(self):
        sql=render(self.stage());self.assertIn('assumed_enrolment_dates',sql);self.assertNotIn('DISABLE ROW LEVEL SECURITY',sql)
        self.assertNotIn('INSERT INTO mass_core.user_scope_grants',sql)
    def history(self,kind,data,student=None):
        path=self.p/'history.xlsx';w=openpyxl.Workbook();s=w.active;s.title='Mpokket - Morning'
        if kind=='attendance':
            s.append(['source']);s.append(['weekday']);s.append(['Name of Child','Grade','Gender','Village Name','Class Category',datetime(2026,7,1)])
        else:
            for i in range(7 if kind=='assessment' else 5):s.append(['header'])
        for r in data:s.append(r)
        w.save(path)
        st=self.stage()['rows'][0];st['donor']='Mpokket'
        return historical({kind:path},[student or st],b'k'*32)
    def test_attendance_uses_scope_and_supporting_details(self):
        p=self.history('attendance',[['Synthetic learner',1,'Girl','B Synthetic','Morning','P']]);self.assertEqual(len(p['accepted']),1)
        p=self.history('attendance',[['Synthetic learner',1,'Girl','Different centre','Morning','P']]);self.assertEqual(len(p['accepted']),0)
    def test_blank_attendance_is_unknown(self):
        p=self.history('attendance',[['Synthetic learner',1,'Girl','B Synthetic','Morning',None]])
        self.assertEqual(p['summary']['attendance']['unrecorded_unknown_cells'],1);self.assertEqual(p['accepted'],[])
    def test_conflicting_duplicate_marks_are_all_held(self):
        p=self.history('attendance',[['Synthetic learner',1,'Girl','B Synthetic','Morning','P'],['Synthetic learner',1,'Girl','B Synthetic','Morning','A']])
        self.assertEqual(p['accepted'],[]);self.assertEqual(len(p['held']),2)
    def test_period_only_requires_explicit_year(self):
        r=['2026-27','Q1','Synthetic learner','Grade 1','Girl',20,None,'B Synthetic']
        p=self.history('assessment',[r]);self.assertEqual(p['accepted'][0]['date'],None)
        r[0]=None;p=self.history('assessment',[r]);self.assertEqual(p['accepted'],[])
    def test_unscoped_morarji_not_matched_by_name(self):
        p=self.history('morarji',[[datetime(2026,10,1),'S1','Synthetic learner','Yes','Yes','Yes','Synthetic School']])
        self.assertEqual(p['accepted'],[]);self.assertEqual(p['held'][0]['reason'],'missing_verified_centre_grade_identity')

if __name__=='__main__':unittest.main()
