"""Offline, review-only XLSX reconciliation. Never connects to a database.
Raw/normalized values and exceptions are written ONLY outside the repository.
No automatic staff/centre identity merges, appointment dates or import SQL.
"""
import argparse, collections, hashlib, json, os, re, unicodedata
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
SHEETS=('Mpokket','Fanuc','McAfee','Janashkati Foundation','Dasar UK')
FIELDS={'centre','region','teacher','officer','shift','donor'}
def normalized(value):
    return re.sub(r'\s+',' ',unicodedata.normalize('NFKC',str(value or ''))).strip()
def reconcile(workbook,mapping):
    book=load_workbook(workbook,read_only=True,data_only=False)
    missing=[s for s in SHEETS if s not in {name.strip() for name in book.sheetnames}]
    if missing:raise ValueError('Required worksheets unavailable; do not substitute another sheet.')
    if len(mapping)!=5 or {name.strip() for name in mapping}!=set(SHEETS):raise ValueError('Explicit verified header mappings are required for all five sheets.')
    digest=hashlib.sha256(Path(workbook).read_bytes()).hexdigest()
    rows=[]
    for ordinal,sheet in enumerate(book.worksheets,1):
        if sheet.title not in mapping:continue
        cfg=mapping[sheet.title];columns={k:v for k,v in cfg['columns'].items() if not(k=='shift' and v is None)};header=int(cfg['header_row'])
        if not {'centre','teacher','officer'}<=set(columns) or not set(columns)<=FIELDS:raise ValueError('Missing or unsupported column mapping.')
        if header<1 or not cfg.get('verified_headers'):raise ValueError('Header row/values must be explicitly verified.')
        from openpyxl.utils.cell import column_index_from_string
        indices={k:column_index_from_string(v) for k,v in columns.items()}
        for field,index in indices.items():
            actual=normalized(sheet.cell(header,index).value)
            if actual!=cfg['verified_headers'].get(field):raise ValueError('Workbook header differs from reviewed mapping; stop.')
        if 'donor' not in columns and not cfg.get('verified_donor_label'):raise ValueError('Donor association requires explicit source-backed mapping; worksheet name alone is insufficient.')
        for n,cells in enumerate(sheet.iter_rows(min_row=header+1,values_only=True),header+1):
            raw={k:cells[i-1] if i<=len(cells) else None for k,i in indices.items()}
            raw={k:None if v is None else str(v) for k,v in raw.items()}
            if not any(normalized(x) for x in raw.values()):continue
            raw.setdefault('shift',None)
            values={k:normalized(v) for k,v in raw.items()};values.setdefault('donor',cfg.get('verified_donor_label',''))
            for field in ('shift','teacher','officer'):
                if not values.get(field):values[field]=None
            issues=[]
            if not values.get('centre'):issues.append('missing_centre')
            if not values.get('donor'):issues.append('missing_donor')
            if values.get('shift') not in ('Morning','Evening',None):issues.append('unverified_shift_value')
            if any(str(v or '').startswith('=') for v in raw.values()):issues.append('formula_value_unverified')
            if any(re.search(r'[,;/&\n]',raw.get(k) or '') for k in ('teacher','officer')):issues.append('possible_multiple_staff_in_cell')
            rows.append(dict(source_sha256=digest,worksheet=sheet.title,worksheet_ordinal=ordinal,row_number=n,source_columns=columns,raw=raw,normalized=values,issues=issues))
    grouped=collections.defaultdict(list);duplicates=collections.defaultdict(list);staff=collections.defaultdict(list)
    for row in rows:
        v=row['normalized'];grouped[v.get('centre','').casefold()].append(row)
        duplicates[(row['worksheet'],tuple(sorted((k,str(x or '').casefold()) for k,x in v.items())))].append(row)
        for role in ('teacher','officer'):
            if v.get(role):staff[(role,v[role].casefold())].append(row)
    for group in duplicates.values():
        if len(group)>1:
            for r in group:r['issues'].append('duplicate_source_row')
    for key,group in grouped.items():
        if not key:continue
        flags=[]
        if len(group)>1:flags.append('centre_identity_review_required')
        if len({r['normalized'].get('donor','').casefold() for r in group})>1:flags.append('possible_multiple_donor_association')
        if len({r['normalized'].get('region','').casefold() for r in group})>1:flags.append('regional_identity_collision')
        if 'masaraguppi' in key and len(group)>1:flags.append('masaraguppi_explicit_review_required')
        for r in group:r['issues'].extend(flags)
    for group in staff.values():
        if len(group)>1:
            for r in group:r['issues'].append('staff_identity_review_required')
    for r in rows:r['issues']=sorted(set(r['issues']))
    # Even unflagged rows need explicit approved canonical IDs before an import.
    summary={'status':'review_required_no_import','source_sha256':digest,'worksheets':{s:sum(r['worksheet']==s for r in rows) for s in mapping},'row_count':len(rows),'flagged_rows':sum(bool(r['issues']) for r in rows),'issue_counts':dict(collections.Counter(i for r in rows for i in r['issues']))}
    book.close();return {'summary':summary,'rows':rows}
def main():
    p=argparse.ArgumentParser();p.add_argument('--workbook',type=Path,required=True);p.add_argument('--mapping',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    out=a.out.resolve()
    if out==ROOT or ROOT in out.parents:raise SystemExit('Private output must be outside the Git/source-archive checkout.')
    result=reconcile(a.workbook,json.loads(a.mapping.read_text()))
    out.mkdir(mode=0o700,parents=True,exist_ok=True);os.chmod(out,0o700)
    f=out/'directory-review.private.json';fd=os.open(f,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
    with os.fdopen(fd,'w') as stream:json.dump(result,stream,ensure_ascii=False,indent=2)
    print('Private review created; no database actions. Rows:',result['summary']['row_count'],'flagged:',result['summary']['flagged_rows'])
if __name__=='__main__':main()
