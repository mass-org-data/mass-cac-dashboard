"""Private-input, structural Cluster Officer reconciliation. Never emits names/notes.

Usage: python scripts/audits/officer_sources.py WORKBOOK --output PRIVATE_JSON
Output must be outside the repository. This tool does not connect to a database.
"""
import argparse
import datetime
import hashlib
import json
from pathlib import Path
import openpyxl


def inspect(path):
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    result = {'filename': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': []}
    for sheet in workbook:
        if 'visit log' not in sheet.title.lower():
            result['sheets'].append({'sheet': sheet.title, 'kind': 'summary', 'individual_visit_evidence': False,
                                    'formula_cells': sum(c.data_type == 'f' for row in sheet for c in row)})
            continue
        headers = {c.column_letter: str(c.value) for c in next(sheet.iter_rows(min_row=6, max_row=6)) if c.value is not None}
        if headers.get('A') != 'Visit Date' or headers.get('B') != 'Center Name' or headers.get('C') != 'Shift':
            raise ValueError('Unknown visit-log header: review privately; no layout guessing')
        rows, examples = [], []
        for row in sheet.iter_rows(min_row=7):
            if len(row) < 4 or row[1].value is None or row[1].data_type == 'f':
                continue
            # Only an explicitly labelled example is excluded. Do not classify an
            # otherwise unsupported record as authentic just because it has a date.
            centre = str(row[1].value).strip()
            if centre.lower().startswith(('e.g', 'example')):
                examples.append(row[1].row)
                continue
            rows.append(row)
        dated = sum(isinstance(row[0].value, (datetime.date, datetime.datetime)) for row in rows)
        result['sheets'].append({
            'sheet': sheet.title, 'kind': 'visit_log', 'headers': headers,
            'example_rows_excluded': examples, 'source_row_numbers': [r[1].row for r in rows],
            'candidate_rows': len(rows), 'dated_candidate_rows': dated,
            'undated_candidate_rows': sum(r[0].value is None for r in rows),
            'unparsed_date_rows': sum(r[0].value is not None and not isinstance(r[0].value, (datetime.date, datetime.datetime)) for r in rows),
            'observed_morning_rows': sum(r[2].value == 'Morning' for r in rows),
            'observed_evening_rows': sum(r[2].value == 'Evening' for r in rows),
            'unknown_shift_rows': sum(r[2].value not in ('Morning', 'Evening') for r in rows),
            'remark_rows': sum(bool(r[3].value) for r in rows),
            'distinct_remark_values': len({r[3].value for r in rows if r[3].value is not None}),
            'officer_identity_column': any('officer' in h.lower() for h in headers.values()),
            'follow_up_column': any('follow' in h.lower() for h in headers.values()),
            'verified_importable_events': 0,
            'reason': 'A source row requires verified date, actual officer identity, centre, shift and authenticity; this audit alone does not establish identity or authenticity.'})
    workbook.close()
    logs = [s for s in result['sheets'] if s['kind'] == 'visit_log']
    result['totals'] = {k: sum(s[k] for s in logs) for k in ['candidate_rows', 'dated_candidate_rows', 'undated_candidate_rows', 'unparsed_date_rows', 'observed_morning_rows', 'observed_evening_rows', 'unknown_shift_rows', 'verified_importable_events']}
    result['totals']['excluded_examples'] = sum(len(s['example_rows_excluded']) for s in logs)
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[2]
    if args.output.resolve().is_relative_to(repo):
        parser.error('Keep source-derived output outside the repository')
    report = inspect(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    args.output.write_text(json.dumps(report, indent=2)+'\n')
    args.output.chmod(0o600)
    print(json.dumps(report['totals']))
