"""Stage immutable source metadata and private review candidates from explicit layouts."""
import json,hashlib,os,re,unicodedata
from collections import Counter
from datetime import date,datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from common import Invalid,canonical,digest,ident,json_value,private_path,read_private,row_hash,text,token
from adapters import FIELDS,observation,iso_day
DIRECTORY={'donor','centre','region','teacher','officer','shift','student_name'}
def stage(workbook,layout,key):
 path=read_private(workbook)
 if len(key)!=32:raise Invalid('Use the existing 32-byte private fingerprint key.')
 if layout.get('format')!='mass-source-layout-v1':raise Invalid('Unsupported verified layout format.')
 kind=layout.get('kind')
 if kind not in {'directory',*FIELDS}:raise Invalid('Unsupported source family.')
 version=token(layout.get('mapping_version'));source_hash=hashlib.sha256(path.read_bytes()).hexdigest()
 file_id=ident('file',source_hash);batch_sheets=[];observations=[]
 book=load_workbook(path,read_only=True,data_only=False)
 configured=layout.get('sheets',{})
 if not configured or not set(configured)<=set(book.sheetnames):raise Invalid('Exact reviewed worksheets are required.')
 if kind=='directory' and set(configured)!={'Mpokket','Fanuc','McAfee','Janashkati Foundation','Dasar UK'}:raise Invalid('Verify all five exact donor worksheets.')
 for ordinal,sheet in enumerate(book.worksheets,1):
  cfg=configured.get(sheet.title);sheet_id=ident('sheet',file_id,ordinal);rows=[]
  header=cfg.get('header_row') if cfg else None
  if cfg:
   if not isinstance(header,int) or header<1:raise Invalid('Explicit header row required.')
   fields=cfg.get('columns',{})
   if not fields or not set(fields)<=(DIRECTORY if kind=='directory' else FIELDS[kind]):raise Invalid('Unsupported or sensitive column selection.')
   if kind=='directory' and not {'centre','teacher','officer','shift'}<=set(fields):raise Invalid('Required directory fields missing.')
   columns={k:column_index_from_string(v) for k,v in fields.items()}
   if len(set(columns.values()))!=len(columns):raise Invalid('Field columns overlap.')
   for field,col in columns.items():
    if json_value(sheet.cell(header,col).value)!=cfg.get('verified_headers',{}).get(field):raise Invalid('Source header drift; stop before staging.')
   attendance={}
   if kind=='attendance':
    for col,expected in cfg.get('attendance_columns',{}).items():
     idx=column_index_from_string(col);actual=sheet.cell(cfg.get('date_header_row',header),idx).value
     if json_value(actual)!=expected:raise Invalid('Attendance date header drift.')
     attendance[idx]=actual
    if not attendance:raise Invalid('Explicit physical attendance/date columns required.')
    attendance_dates=Counter(iso_day(v) for v in attendance.values())
  for number,values in enumerate(sheet.iter_rows(values_only=True),1):
   native_values=list(values)
   values=[json_value(v) for v in native_values]
   row_id=ident('row',sheet_id,number)
   candidate=bool(cfg and number>header and any(v not in (None,'') for v in values))
   row={'id':row_id,'row_number':number,'content_hash':row_hash(values,key),'classification':'candidate' if candidate else 'scaffold'}
   rows.append(row)
   if not candidate:continue
   raw={k:native_values[c-1] if c<=len(native_values) else None for k,c in columns.items()}
   reference={'source_row_id':row_id,'worksheet_ordinal':ordinal,'row_number':number,'source_columns':columns}
   if kind=='directory':
    # Preserve selected original strings; normalization is comparison-only.
    raw={k:json_value(v) for k,v in raw.items()}
    normalized={k:re.sub(r'\s+',' ',unicodedata.normalize('NFKC',str(v or ''))).strip() for k,v in raw.items()};issues=[]
    if not normalized.get('donor'):
     if not cfg.get('verified_donor_label'):issues.append('missing_verified_donor')
     else:normalized['donor']=cfg['verified_donor_label']
    if not normalized.get('centre'):issues.append('missing_centre')
    for field in ('teacher','officer'):
     if not normalized.get(field):issues.append('missing_'+field)
     if re.search(r'[,;/&\n]',str(raw.get(field) or '')):issues.append('possible_multiple_staff_in_cell')
     if re.search(r'@|(?:\+?\d[\s().-]*){7,}',str(raw.get(field) or '')):issues.append('contact_detail_in_staff_cell')
    if normalized.get('shift') not in ('','Morning','Evening'):issues.append('unverified_shift')
    if any(str(v or '').startswith('=') for v in raw.values()):issues.append('formula_value_unverified')
    observations.append({**reference,'raw':raw,'values':normalized,'issues':issues,'operational_import_ready':False})
   elif kind=='attendance':
    for col,on in attendance.items():
     code=values[col-1] if col<=len(values) else None
     result=observation(kind,raw,{'date':on,'code':code})
     if iso_day(on) and attendance_dates[iso_day(on)]>1:result['issues'].append('duplicate_attendance_date_header')
     observations.append({**reference,'source_slot':col,**result})
   else:observations.append({**reference,**observation(kind,raw)})
  batch_sheets.append({'id':sheet_id,'ordinal':ordinal,'original_name':sheet.title,'classification':'input' if cfg else 'unclassified','header_version':int(cfg.get('header_version',1)) if cfg else 1,'rows':rows})
 book.close()
 if kind=='directory':flag_collisions(observations)
 source={'id':file_id,'sha256':source_hash,'storage_ref':ident('storage',source_hash),'original_file_id':layout.get('original_file_id')}
 if source['original_file_id'] is not None:
  if not re.fullmatch(r'[A-Za-z0-9_-]{1,128}',source['original_file_id']):raise Invalid('Invalid known original file identifier.')
 return {'format':'mass-import-plan-v1','kind':kind,'mapping_version':version,'extractor_version':'mass-import-v1','fingerprint_key_id':hashlib.sha256(key).hexdigest(),'source':source,'sheets':batch_sheets,'observations':observations}

def flag_collisions(items):
 from collections import defaultdict
 groups=defaultdict(list);staff=defaultdict(list);duplicates=defaultdict(list)
 for o in items:
  v=o['values'];groups[v.get('centre','').casefold()].append(o)
  duplicates[(o['worksheet_ordinal'],digest(v))].append(o)
  for role in ('teacher','officer','student_name'):
   if v.get(role):staff[(role,v[role].casefold())].append(o)
 for key,group in groups.items():
  if not key:continue
  flags=[]
  if len(group)>1:flags.append('centre_identity_review')
  if len({x['values'].get('donor','').casefold() for x in group})>1:flags.append('possible_multiple_donors')
  if len({x['values'].get('region','').casefold() for x in group})>1:flags.append('regional_identity_collision')
  if 'masaraguppi' in key and len(group)>1:flags.append('masaraguppi_review')
  for o in group:o['issues'].extend(flags)
 for (role,_),group in staff.items():
  if len(group)>1:
   for o in group:o['issues'].append('student_name_identity_review' if role=='student_name' else 'staff_identity_review')
 for group in duplicates.values():
  if len(group)>1:
   for o in group:o['issues'].append('duplicate_source_observation')
 for o in items:o['issues']=sorted(set(o['issues']))
